"""Tests for tracker/excel_import.py and service.import_excel — reading earlier Excel exports back in."""

import io
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

from tracker import service
from tracker.database import TimesheetDB
from tracker.excel_import import ExcelImportError, WeekNeeded, parse_workbook, week_from_filename
from tracker.models import Goal, GoalOutcome, Reflection, Subject, TimeEntry
from tracker.week_report import build_week_report_xlsx, report_filename

MON = date(2026, 4, 20)  # a Monday


@pytest.fixture
def src(tmp_path):
    """A database with one week of data, to export from."""
    with TimesheetDB(str(tmp_path / "src.db")) as d:
        yield d


@pytest.fixture
def dst(tmp_path):
    """An empty database, to import into."""
    with TimesheetDB(str(tmp_path / "dst.db")) as d:
        yield d


def fill(db):
    a = db.add_subject(Subject(name="Programming", low_level_label="python", high_level_label="Work"))
    b = db.add_subject(Subject(name="Reading", low_level_label="novels", high_level_label="Personal"))
    db.add_entry(TimeEntry(date=MON, subject_id=a.id, duration_hours=2.5, notes="parser"))
    db.add_entry(TimeEntry(date=MON, subject_id=a.id, duration_hours=1.0))
    db.add_entry(TimeEntry(date=MON + timedelta(days=3), subject_id=b.id, duration_hours=0.75, notes="ch. 4"))
    db.upsert_reflection(Reflection(week_start=MON, strengths="focus", weaknesses="late starts", next_week_plan="mornings"))
    g = db.add_goal(Goal(week_start=MON, description="Code 3h", target_hours=3, subject_id=a.id))
    db.add_goal(Goal(week_start=MON, description="Rest"))
    db.upsert_goal_outcome(GoalOutcome(goal_id=g.id, met=1, notes="done"))
    return a, b


def report_bytes(db, week=MON):
    goals = db.get_goals_for_week(week)
    outcomes = {g.id: o for g in goals if (o := db.get_outcome_for_goal(g.id)) is not None}
    return build_week_report_xlsx(week, db.get_entries_for_week(week), db.get_reflection(week), goals, outcomes)


def table_view(db, week=MON):
    return sorted((r["name"], r["low"], r["high"], tuple(r["days"])) for r in service.week_table(db, week)["rows"])


# ------------------------------------------------------------------ #
# Reading files
# ------------------------------------------------------------------ #

class TestParse:
    def test_week_from_filename(self):
        assert week_from_filename("time_report_2026-04-22.xlsx") == MON  # snapped to Monday
        assert week_from_filename("my hours.xlsx") is None

    def test_report(self, src):
        fill(src)
        wb = parse_workbook(report_bytes(src), "whatever.xlsx")  # no date in the name: the entries' dates are used
        assert (wb.kind, wb.week, wb.week_source) == ("report", MON, "entries")
        assert sorted((e.date, e.name, e.hours, e.notes) for e in wb.entries) == [
            (MON, "Programming", 1.0, ""), (MON, "Programming", 2.5, "parser"),
            (MON + timedelta(days=3), "Reading", 0.75, "ch. 4")]
        assert wb.reflection == {"strengths": "focus", "weaknesses": "late starts", "plan": "mornings"}
        assert [(g["description"], g["met"], g["target_hours"]) for g in wb.goals] == [("Code 3h", 1, 3), ("Rest", None, None)]

    def test_table_download_uses_name(self, src):
        fill(src)
        wb = parse_workbook(service.week_table_xlsx(src, MON), "timesheet_2026-04-20.xlsx")
        assert (wb.kind, wb.week, wb.week_source) == ("table", MON, "file name")
        # The table has day totals only: 2.5 + 1.0 become one 3.5 h cell.
        assert sorted((e.date, e.name, e.hours) for e in wb.entries) == [
            (MON, "Programming", 3.5), (MON + timedelta(days=3), "Reading", 0.75)]

    def test_table_without_date_needs_week(self, src):
        fill(src)
        data = service.week_table_xlsx(src, MON)
        with pytest.raises(WeekNeeded):
            parse_workbook(data, "hours.xlsx")
        wb = parse_workbook(data, "hours.xlsx", week=MON + timedelta(days=2))
        assert (wb.week, wb.week_source) == (MON, "chosen")

    def test_chosen_week_never_overrides_a_dated_file(self, src):
        fill(src)
        wb = parse_workbook(service.week_table_xlsx(src, MON), "timesheet_2026-04-20.xlsx", week=date(2026, 1, 5))
        assert (wb.week, wb.week_source) == (MON, "file name")

    def test_name_and_dates_disagree(self, src):
        fill(src)
        wb = parse_workbook(report_bytes(src), "time_report_2026-01-05.xlsx")
        assert {e.date for e in wb.entries} >= {MON}  # entries keep their own dates
        assert any("entries are dated" in w for w in wb.warnings)

    def test_not_excel(self):
        with pytest.raises(ExcelImportError, match="not an Excel"):
            parse_workbook(b"hello", "x.xlsx")

    def test_foreign_workbook(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Name", "Amount"])
        buf = io.BytesIO(); wb.save(buf)
        with pytest.raises(ExcelImportError, match="no Time Tracker data"):
            parse_workbook(buf.getvalue(), "budget_2026-04-20.xlsx")

    def test_bad_hours_skipped_with_note(self):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "Entries"
        ws.append(["Date", "Subject", "Low Label", "High Label", "Hours", "Notes"])
        ws.append([MON, "A", "a", "A", 30, ""])
        ws.append([MON, "A", "a", "A", 2, ""])
        buf = io.BytesIO(); wb.save(buf)
        parsed = parse_workbook(buf.getvalue(), "x.xlsx")
        assert [e.hours for e in parsed.entries] == [2]
        assert any("not between 0 and 24" in w for w in parsed.warnings)


# ------------------------------------------------------------------ #
# Importing
# ------------------------------------------------------------------ #

class TestImport:
    def test_report_round_trip(self, src, dst):
        fill(src)
        res = service.import_excel(dst, [(report_filename(MON), report_bytes(src))])
        assert res["files"][0]["ok"]
        assert table_view(dst) == table_view(src)
        assert dst.get_reflection(MON).next_week_plan == "mornings"
        goals = {g.description: g for g in dst.get_goals_for_week(MON)}
        assert set(goals) == {"Code 3h", "Rest"}
        assert goals["Code 3h"].subject_name == "Programming"  # linked again, so Goal Review can auto-evaluate
        assert dst.get_outcome_for_goal(goals["Code 3h"].id).met == 1
        assert dst.get_outcome_for_goal(goals["Rest"].id) is None
        assert res["weeks"] == [{"week_start": "2026-04-20", "hours_before": 0, "hours_after": 4.25}]
        # Notes survive (the Entries sheet carries them).
        assert {e.notes for e in dst.get_entries_for_week(MON)} == {"parser", "", "ch. 4"}

    def test_second_import_changes_nothing(self, src, dst):
        fill(src)
        # Trailing spaces must survive the round trip, or the second import sees a change.
        src.upsert_reflection(Reflection(week_start=MON, strengths="focus ", next_week_plan="2 more hours. "))
        files = [(report_filename(MON), report_bytes(src))]
        service.import_excel(dst, files)
        res = service.import_excel(dst, files)
        f = res["files"][0]
        assert f["entries"]["new"] == 0 and f["entries"]["different"] == 0
        assert f["reflection"] == "same" and f["goals"] == {"new": 0, "existing": 2}
        res = service.import_excel(dst, files, mode="replace")
        assert res["files"][0]["reflection"] == "same"
        assert len(dst.get_entries_for_week(MON)) == 3
        assert len(dst.get_goals_for_week(MON)) == 2

    def test_report_and_table_of_same_week_together(self, src, dst):
        fill(src)
        res = service.import_excel(dst, [(report_filename(MON), report_bytes(src)),
                                         ("timesheet_2026-04-20.xlsx", service.week_table_xlsx(src, MON))])
        assert res["files"][1]["entries"] == {"new": 0, "same": 2, "different": 0}
        assert table_view(dst) == table_view(src)

    def test_keep_vs_replace(self, src, dst):
        a, _ = fill(src)
        mine = dst.add_subject(Subject(name="Programming", low_level_label="python", high_level_label="Work"))
        dst.add_entry(TimeEntry(date=MON, subject_id=mine.id, duration_hours=1.0, notes="mine"))
        dst.upsert_reflection(Reflection(week_start=MON, strengths="my own"))
        files = [(report_filename(MON), report_bytes(src))]

        res = service.import_excel(dst, files, mode="keep")
        assert res["files"][0]["entries"]["different"] == 1
        assert res["files"][0]["reflection"] == "kept"
        monday = [e for e in dst.get_entries_for_week(MON) if e.date == MON]
        assert [(e.duration_hours, e.notes) for e in monday] == [(1.0, "mine")]
        assert dst.get_reflection(MON).strengths == "my own"

        res = service.import_excel(dst, files, mode="replace")
        assert res["files"][0]["reflection"] == "replaced"
        monday = [e for e in dst.get_entries_for_week(MON) if e.date == MON]
        assert sorted(e.duration_hours for e in monday) == [1.0, 2.5]  # the file's two entries
        assert dst.get_reflection(MON).strengths == "focus"

    def test_unknown_label_matches_empty_label(self, dst):
        s = dst.add_subject(Subject(name="jkj", low_level_label="", high_level_label=""))
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "Entries"
        ws.append(["Date", "Subject", "Low Label", "High Label", "Hours", "Notes"])
        ws.append([MON, "jkj", "Unknown", "Unknown", 1, ""])
        buf = io.BytesIO(); wb.save(buf)
        service.import_excel(dst, [("x.xlsx", buf.getvalue())])
        assert [e.subject_id for e in dst.get_entries_for_week(MON)] == [s.id]
        assert len(dst.get_all_subjects()) == 1

    def test_bad_file_does_not_stop_others(self, src, dst):
        fill(src)
        res = service.import_excel(dst, [("broken.xlsx", b"nope"), ("hours.xlsx", service.week_table_xlsx(src, MON)),
                                         (report_filename(MON), report_bytes(src))])
        ok = [f["ok"] for f in res["files"]]
        assert ok == [False, False, True]
        assert res["files"][1]["needs_week"]
        assert table_view(dst) == table_view(src)

    def test_chosen_week(self, src, dst):
        fill(src)
        res = service.import_excel(dst, [("hours.xlsx", service.week_table_xlsx(src, MON))], weeks={"0": "2026-04-22"})
        assert res["files"][0]["ok"] and res["files"][0]["week_source"] == "chosen"
        assert len(dst.get_entries_for_week(MON)) == 2

    def test_removed_row_comes_back(self, src, dst):
        fill(src)
        subj = dst.add_subject(Subject(name="Reading", low_level_label="novels", high_level_label="Personal"))
        service.remove_row(dst, MON, subj.id)
        service.import_excel(dst, [(report_filename(MON), report_bytes(src))])
        assert subj.id in {r["id"] for r in service.week_table(dst, MON)["rows"]}

    def test_preview_changes_nothing(self, src, tmp_path):
        fill(src)
        target = tmp_path / "target.db"
        with TimesheetDB(str(target)):
            pass
        res = service.preview_import(target, [(report_filename(MON), report_bytes(src))])
        assert res["files"][0]["entries"]["new"] == 2
        assert res["weeks"][0]["hours_after"] == 4.25
        with TimesheetDB(str(target)) as d:
            assert d.get_entries_for_week(MON) == [] and d.get_all_subjects() == []

    def test_preview_without_database_yet(self, src, tmp_path):
        fill(src)
        res = service.preview_import(tmp_path / "none.db", [(report_filename(MON), report_bytes(src))])
        assert res["files"][0]["ok"]
        assert not (tmp_path / "none.db").exists()

    def test_bad_mode(self, dst):
        with pytest.raises(service.UserError):
            service.import_excel(dst, [], mode="merge")
