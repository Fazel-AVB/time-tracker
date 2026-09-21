"""Unit tests for tracker/week_report.py — pending weeks and the report workbook."""

import io
from datetime import date, timedelta

import openpyxl

from tracker.models import Goal, GoalOutcome, Reflection, TimeEntry
from tracker.week_report import build_week_report_xlsx, pending_export_weeks, report_filename

MON = date(2026, 9, 14)          # a Monday
NEXT = MON + timedelta(weeks=1)  # 2026-09-21


def _entry(day, hours, name="Reading", sid=1, notes=""):
    return TimeEntry(date=MON + timedelta(days=day), subject_id=sid, duration_hours=hours, notes=notes,
                     subject_name=name, low_level_label="leisure", high_level_label="Personal")


class TestPendingExportWeeks:
    def test_finished_week_with_entries_is_pending(self):
        assert pending_export_weeks(NEXT, since=MON, weeks_with_entries={MON}, decided=set()) == [MON]

    def test_current_week_not_pending_even_on_sunday(self):
        sunday = MON + timedelta(days=6)
        assert pending_export_weeks(sunday, since=MON, weeks_with_entries={MON}, decided=set()) == []

    def test_decided_week_not_pending(self):
        assert pending_export_weeks(NEXT, MON, {MON}, decided={MON}) == []

    def test_weeks_before_since_ignored(self):
        old = MON - timedelta(weeks=3)
        assert pending_export_weeks(NEXT, since=MON, weeks_with_entries={old, MON}, decided=set()) == [MON]

    def test_week_without_entries_not_pending(self):
        assert pending_export_weeks(NEXT, MON, set(), set()) == []

    def test_oldest_first(self):
        w0 = MON - timedelta(weeks=1)
        assert pending_export_weeks(NEXT, w0, {MON, w0}, set()) == [w0, MON]


class TestBuildWeekReportXlsx:
    def _open(self, data):
        return openpyxl.load_workbook(io.BytesIO(data))

    def test_filename(self):
        assert report_filename(MON) == "time_report_2026-09-14.xlsx"

    def test_core_sheets_and_daily_total(self):
        wb = self._open(build_week_report_xlsx(MON, [_entry(0, 1.5), _entry(1, 2.0)]))
        assert wb.sheetnames == ["Weekly Table", "By High Label", "By Low Label", "By Subject", "Entries"]
        rows = list(wb["Weekly Table"].values)
        header, last = rows[0], rows[-1]
        assert last[0] == "DAILY TOTAL"
        assert last[header.index("Total")] == 3.5
        assert wb["Entries"].max_row == 3  # header + 2 entries

    def test_reflection_and_goals_sheets(self):
        goal = Goal(id=7, week_start=MON, description="Read 4h", target_hours=4.0, subject_id=1,
                    subject_name="Reading")
        data = build_week_report_xlsx(
            MON, [_entry(0, 3.0)],
            reflection=Reflection(week_start=MON, strengths="focus"),
            goals=[goal], outcomes={7: GoalOutcome(goal_id=7, met=2, notes="close")},
        )
        wb = self._open(data)
        goals = list(wb["Goals"].values)
        assert goals[1][goals[0].index("Logged hours")] == 3.0
        assert goals[1][goals[0].index("Outcome")] == "Partial"
        assert list(wb["Reflection"].values)[1] == ("Strengths", "focus")

    def test_empty_week_still_builds(self):
        wb = self._open(build_week_report_xlsx(MON, []))
        assert "Weekly Table" in wb.sheetnames
